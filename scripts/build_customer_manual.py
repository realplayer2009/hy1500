#!/usr/bin/env python3
"""将客户使用说明书 Markdown 和本地插图生成 Word。

用法：python3 scripts/build_customer_manual.py [说明书.md]
依赖：python-docx、Pillow。PDF 使用 LibreOffice 从同一 Word 导出。
"""

import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def text_runs(paragraph, value):
    value = re.sub(r"(\d) (℃|%RH|V|GB|秒|分钟)", r"\1" + "\u00a0" + r"\2", value)
    for index, part in enumerate(re.split(r"\*\*(.*?)\*\*", value)):
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1


def add_table(document, rows):
    widths = [3.2, 11.0] if len(rows[0]) == 2 else [14.2 / len(rows[0])] * len(rows[0])
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + side)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 90), ("bottom", 90), ("left", 110), ("right", 110)):
        element = OxmlElement("w:" + side)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    table._tbl.tblPr.append(margins)
    for index, values in enumerate(rows):
        row = table.rows[0] if index == 0 else table.add_row()
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for column_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EEEEEE" if index == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(14)
            if rows[0][0] == "操作" and column_index == 2:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.bold = index == 0
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = Pt(2)
    paragraph.add_run().font.size = Pt(2)


source = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "docs/客户使用说明书.md"

# 图示按开发资料绘制板端关系；不替代整箱端子图和扩展板接线资料。
font_path = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
diagram_font = ImageFont.truetype(font_path, 42)
diagram_small = ImageFont.truetype(font_path, 36)
diagram_large = ImageFont.truetype(font_path, 48)


def label(draw, xy, value, font=diagram_font):
    draw.text(xy, value, font=font, fill="black", anchor="mm")


def arrow(draw, start, end):
    draw.line((start, end), fill="black", width=4)
    x, y = end
    draw.line(((x - 14, y - 10), end, (x - 14, y + 10)), fill="black", width=4)


def heater(draw, x, y, title):
    draw.rounded_rectangle((x, y, x + 140, y + 130), radius=10, outline="black", width=4)
    for offset in (30, 55, 80, 105):
        draw.line(((x + offset, y + 25), (x + offset, y + 105)), fill="black", width=3)
    label(draw, (x + 70, y + 170), title)


image_dir = source.parent / "客户使用说明书插图"
image_dir.mkdir(exist_ok=True)
# 分体式控制箱、加热器与探头的关系，不标注未确认的实物尺寸。
canvas = Image.new("RGB", (1600, 820), "white")
draw = ImageDraw.Draw(canvas)
draw.rounded_rectangle((490, 50, 1110, 350), radius=18, outline="black", width=5)
draw.rectangle((550, 95, 1050, 240), outline="black", width=3)
label(draw, (800, 165), "加热器控制系统", diagram_large)
for x in (655, 800, 945):
    draw.ellipse((x - 20, 270, x + 20, 310), outline="black", width=3)
label(draw, (800, 405), "控制箱")
draw.line(((800, 445), (800, 505), (310, 505), (1290, 505)), fill="black", width=4)
for x, title in ((240, "加热器 A"), (730, "加热器 B"), (1220, "加热器 C")):
    draw.line(((x + 70, 505), (x + 70, 565)), fill="black", width=4)
    heater(draw, x, 565, title)
canvas.save(image_dir / "system.png")

canvas = Image.new("RGB", (1600, 1000), "white")
draw = ImageDraw.Draw(canvas)
draw.rectangle((550, 70, 1510, 850), outline="black", width=5)
label(draw, (1030, 120), "变压器箱体")
draw.rectangle((70, 210, 410, 600), outline="black", width=5)
draw.rectangle((110, 265, 370, 405), outline="black", width=3)
label(draw, (240, 335), "控制屏")
label(draw, (240, 530), "控制箱")
label(draw, (240, 675), "便于观察操作", diagram_small)
draw.line(((410, 510), (500, 510), (500, 780), (630, 780)), fill="black", width=4)
label(draw, (360, 800), "连接线缆", diagram_small)
for x, title in ((660, "加热器 A"), (970, "加热器 B"), (1280, "加热器 C")):
    heater(draw, x, 555, title)
    draw.line(((x + 70, 490), (x + 70, 400)), fill="black", width=3)
    draw.line(((x + 56, 425), (x + 70, 400), (x + 84, 425)), fill="black", width=3)
label(draw, (1040, 350), "保持进出风通畅")
draw.ellipse((665, 195, 725, 255), outline="black", width=3)
label(draw, (1010, 225), "温湿度探头避开出风口")
label(draw, (800, 930), "位置关系示意  实际固定与走线由供方确定", diagram_small)
canvas.save(image_dir / "installation.png")

canvas = Image.new("RGB", (1600, 1120), "white")
draw = ImageDraw.Draw(canvas)
draw.rectangle((70, 100, 610, 1080), outline="black", width=4)
label(draw, (340, 155), "主板接口")
for y, terminals, target, detail in (
    (280, "A1+ / B1−", "控制屏 RS485", "按 + / − 极性对应"),
    (470, "A2+ / B2−", "温湿度探头 RS485", "探头供电按额定值与随箱图连接"),
    (660, "A3+ / B3−", "PT100 采集模块 RS485", "模块再接 PT100 探头"),
    (880, "+V / 0V", "板卡直流电源", "正极接 +V，负极接 0V"),
):
    label(draw, (340, y), terminals)
    draw.rectangle((880, y - 60, 1540, y + 60), outline="black", width=3)
    label(draw, (1210, y), target)
    arrow(draw, (610, y), (880, y))
    label(draw, (1195, y + 105), detail, diagram_small)
label(draw, (340, 1020), "板端标识供核对", diagram_small)
canvas.save(image_dir / "wiring.png")

canvas = Image.new("RGB", (1600, 850), "white")
draw = ImageDraw.Draw(canvas)
label(draw, (800, 45), "加热器控制回路")
for y, title in ((150, "A"), (270, "B"), (390, "C")):
    for left, right, value in ((60, 440, "OT 控制信号"), (610, 1050, "外接继电器回路"), (1240, 1540, "加热器 " + title)):
        draw.rectangle((left, y - 45, right, y + 45), outline="black", width=3)
        label(draw, ((left + right) // 2, y), value, diagram_small)
    arrow(draw, (440, y), (610, y))
    arrow(draw, (1050, y), (1240, y))
label(draw, (800, 480), "OT 对地导通控制  负载由独立供电及保护回路供电", diagram_small)
label(draw, (800, 565), "高压检测接口")
for y, value, port in ((655, "适配后的模拟检测信号", "Vin / GND"), (775, "按验收接法的数字检测信号", "输入 / GND")):
    draw.rectangle((60, y - 42, 800, y + 42), outline="black", width=3)
    label(draw, (430, y), value, diagram_small)
    draw.rectangle((1060, y - 42, 1540, y + 42), outline="black", width=3)
    label(draw, (1300, y), port)
    arrow(draw, (800, y), (1060, y))
canvas.save(image_dir / "heater_wiring.png")

document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(17), Cm(24)
section.left_margin = section.right_margin = Cm(1.4)
section.top_margin = section.bottom_margin = Cm(1.2)
section.footer_distance = Cm(0.65)
for grid in section._sectPr.findall(qn("w:docGrid")):
    section._sectPr.remove(grid)
for border in document.styles.element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
for name, size in (("Normal", 11), ("Title", 22), ("Subtitle", 10), ("Heading 1", 17), ("Heading 2", 12), ("List Bullet", 11), ("Caption", 9)):
    style = document.styles[name]
    style.font.name = "Noto Sans CJK SC"
    style.font.size = Pt(size)
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string("000000")
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = Pt(17)
    style.paragraph_format.widow_control = True
    style.paragraph_format.keep_together = True
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    style.element.get_or_add_pPr().append(snap)
    for color in style.element.xpath(".//w:color"):
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn("w:" + attr), None)
    if name == "Title":
        style.font.bold = True
        style.paragraph_format.line_spacing = Pt(30)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(16)
    elif name.startswith("Heading"):
        style.font.bold = True
        style.paragraph_format.line_spacing = Pt(26 if name == "Heading 1" else 18)
        style.paragraph_format.space_before = Pt(8 if name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    elif name in ("Subtitle", "Caption"):
        style.font.color.rgb = RGBColor.from_string("555555")
        style.paragraph_format.line_spacing = Pt(13)
    elif name == "List Bullet":
        style.paragraph_format.left_indent = Cm(0.45)
        style.paragraph_format.first_line_indent = Cm(-0.35)
        style.paragraph_format.keep_with_next = False

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("裕泰通加热器控制系统    ")
run.font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

lines = source.read_text(encoding="utf-8").splitlines()
index = 0
page_break_pending = False
while index < len(lines):
    line = lines[index].strip()
    index += 1
    if not line:
        continue
    if line == "<!-- pagebreak -->":
        page_break_pending = True
        continue
    if line.startswith("|"):
        rows = []
        while True:
            row = [value.strip() for value in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", value) for value in row):
                rows.append(row)
            if index == len(lines) or not lines[index].strip().startswith("|"):
                break
            line = lines[index].strip()
            index += 1
        add_table(document, rows)
        continue
    image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
    if image_match:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(3)
        picture = paragraph.add_run().add_picture(str(source.parent / image_match[2]), width=Cm(14.2))
        picture._inline.docPr.set("descr", image_match[1])
        continue
    if line.startswith("# "):
        paragraph = document.add_paragraph(line[2:], "Title")
    elif line.startswith("## "):
        paragraph = document.add_paragraph(style="Heading 1")
        number = re.match(r"^(\d+) (.*)", line[3:])
        if number:
            run = paragraph.add_run(number[1] + "  ")
            run.font.size = Pt(20)
            paragraph.add_run(number[2])
        else:
            paragraph.add_run(line[3:])
    elif line.startswith("### "):
        paragraph = document.add_paragraph(line[4:], "Heading 2")
    elif line == "客户使用说明书":
        paragraph = document.add_paragraph(line, "Subtitle")
    elif line.startswith("图 "):
        paragraph = document.add_paragraph(line, "Caption")
    else:
        paragraph = document.add_paragraph(style="List Bullet" if line.startswith("- ") else "Normal")
        step = re.match(r"^(\d+)\. (.*)", line)
        if step:
            paragraph.paragraph_format.left_indent = Cm(0.85)
            paragraph.paragraph_format.first_line_indent = Cm(-0.85)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.85))
            run = paragraph.add_run(step[1] + "\t")
            run.bold = True
            run.font.size = Pt(11)
            text_runs(paragraph, step[2])
            paragraph.paragraph_format.line_spacing = Pt(16)
            paragraph.paragraph_format.space_after = Pt(7)
        else:
            text_runs(paragraph, line[2:] if line.startswith("- ") else line)
    if page_break_pending:
        paragraph.paragraph_format.page_break_before = True
        page_break_pending = False

document.core_properties.title = "裕泰通加热器控制系统客户使用说明书"
document.core_properties.subject = "接线与基本使用"
document.core_properties.author = ""
document.core_properties.comments = ""
output = source.with_suffix(".docx")
document.save(output)
print(output)
